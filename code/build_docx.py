#!/usr/bin/env python3
"""Assemble the manuscript draft markdown files into a single Word document,
rendering the equations as native Word equation objects (OMML)."""

from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import parse_xml
from docx.shared import Pt, RGBColor, Inches

BASE = Path(__file__).resolve().parent.parent  # hepatotoxicity dir

TITLE = ("How much does assumed enzyme zonation matter for predicted "
         "acetaminophen hepatotoxicity? A stability analysis against measured "
         "human profiles")
SUBTITLE = ""

FILES = [
    "manuscript/front_matter.md",
    "manuscript/Introduction_draft.md",
    "manuscript/Methods_Results_draft.md",
    "manuscript/Discussion_draft.md",
    "manuscript/References_draft.md",
    "manuscript/Supplementary_Methods.md",
]

INLINE = re.compile(r"(\*\*.+?\*\*|\*[^*\s][^*]*\*)")

M_NS = "http://schemas.openxmlformats.org/officeDocument/2006/math"


# --- OMML equation builder -------------------------------------------------
def _esc(t: str) -> str:
    return t.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")


def _r(t: str) -> str:  # plain run
    return f'<m:r><m:t xml:space="preserve">{_esc(t)}</m:t></m:r>'


def _frac(num: str, den: str) -> str:
    return f"<m:f><m:num>{num}</m:num><m:den>{den}</m:den></m:f>"


def _ssub(base: str, sub: str) -> str:
    return f"<m:sSub><m:e>{base}</m:e><m:sub>{sub}</m:sub></m:sSub>"


def _ssup(base: str, sup: str) -> str:
    return f"<m:sSup><m:e>{base}</m:e><m:sup>{sup}</m:sup></m:sSup>"


def _k(sub): return _ssub(_r("k"), _r(sub))
def _b(sub): return _ssub(_r("b"), _r(sub))
def _d(sub): return _ssub(_r("d"), _r(sub))


# keyed by the leading token so the mapping survives minor line edits
EQUATIONS = [
    ("dP/dt", (_frac(_r("dP"), _r("dt")) + _r(" = −") + _k("S") + _r("·S·P")
               + _r(" − ") + _k("G") + _r("·P") + _r(" − ") + _k("450") + _r("·P")
               + _r(" + ") + _k("N") + _r("·N"))),
    ("dS/dt", (_frac(_r("dS"), _r("dt")) + _r(" = −") + _k("S") + _r("·S·P")
               + _r(" + ") + _b("S") + _r(" − ") + _d("S") + _r("·S"))),
    ("dN/dt", (_frac(_r("dN"), _r("dt")) + _r(" = ") + _k("450") + _r("·P")
               + _r(" − ") + _k("N") + _r("·N") + _r(" − ") + _k("GSH") + _r("·N·G")
               + _r(" − ") + _k("PSH") + _r("·N"))),
    ("dG/dt", (_frac(_r("dG"), _r("dt")) + _r(" = −") + _k("GSH") + _r("·N·G")
               + _r(" + ") + _b("G") + _r(" − ") + _d("G") + _r("·G"))),
    ("dC/dt", (_frac(_r("dC"), _r("dt")) + _r(" = ") + _k("PSH") + _r("·N")
               + _r(" − ") + _k("clear") + _r("·C"))),
    ("m(x)", (_r("m(x) = ") + _ssup(_r("F"), _r("x − 1/2"))
              + _r(",   x = ") + _frac(_r("i − 1"), _r("15")) + _r(","))),
]


def equation_math(line: str) -> str:
    s = line.strip()
    for prefix, math in EQUATIONS:
        if s.startswith(prefix):
            return math
    return _r(s)  # unknown line -> plain-text equation (still a native object)


def add_equation(doc, math_xml: str):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(2)
    omathpara = parse_xml(
        f'<m:oMathPara xmlns:m="{M_NS}"><m:oMath>{math_xml}</m:oMath></m:oMathPara>'
    )
    p._p.append(omathpara)


_fig_num = 0


def add_figure(doc, fname: str, caption: str):
    global _fig_num
    _fig_num += 1
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(2)
    run = p.add_run()
    run.add_picture(str(BASE / "figures" / fname), width=Inches(6.1))
    cp = doc.add_paragraph()
    cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cp.paragraph_format.space_after = Pt(8)
    cr = cp.add_run(f"Fig. {_fig_num}. ")
    cr.bold = True
    cr.font.size = Pt(9)
    c2 = cp.add_run(caption)
    c2.font.size = Pt(9)


# --- markdown parsing ------------------------------------------------------
def add_rich(paragraph, text, base_bold=False):
    """Add runs to a paragraph, parsing **bold** and *italic* inline."""
    for part in INLINE.split(text):
        if not part:
            continue
        if part.startswith("**") and part.endswith("**") and len(part) >= 4:
            r = paragraph.add_run(part[2:-2])
            r.bold = True
            r.italic = base_bold
        elif part.startswith("*") and part.endswith("*") and len(part) >= 3:
            r = paragraph.add_run(part[1:-1])
            r.italic = True
            r.bold = base_bold
        else:
            r = paragraph.add_run(part)
            r.bold = base_bold
            r.italic = base_bold


def split_row(line: str):
    return [c.strip() for c in line.strip().strip("|").split("|")]


def is_separator_row(cells):
    return len(cells) > 0 and all(re.match(r"^:?-{2,}:?$", c) for c in cells if c != "")


def add_table(doc, rows):
    header, body = rows[0], rows[1:]
    ncols = len(header)
    table = doc.add_table(rows=1, cols=ncols)
    table.style = "Table Grid"
    for j, h in enumerate(header):
        p = table.rows[0].cells[j].paragraphs[0]
        add_rich(p, h)
        for r in p.runs:
            r.bold = True
    for row in body:
        cells = table.add_row().cells
        for j in range(ncols):
            val = row[j] if j < len(row) else ""
            add_rich(cells[j].paragraphs[0], val)


def process_file(doc, path):
    text = path.read_text(encoding="utf-8")
    lines = text.splitlines()

    in_code = False
    code_buf = []
    para_buf = []
    table_buf = []

    def flush_para():
        nonlocal para_buf
        if para_buf:
            p = doc.add_paragraph()
            add_rich(p, " ".join(para_buf))
            para_buf = []

    def flush_table():
        nonlocal table_buf
        if table_buf:
            rows = [r for r in table_buf if not is_separator_row(r)]
            if rows:
                add_table(doc, rows)
            table_buf = []

    for line in lines:
        line = line.rstrip()

        if line.startswith("# ") and not line.startswith("## "):
            continue
        if line.startswith("**Status:"):
            continue
        if line.strip() == "---":
            continue

        # figure directive: [[FIG <file> | <caption>]]
        m = re.match(r"^\[\[FIG\s+(.+?)\s*\|\s*(.+?)\]\]\s*$", line)
        if m:
            flush_para()
            flush_table()
            add_figure(doc, m.group(1).strip(), m.group(2).strip())
            continue

        # code fences -> equations
        if line.strip().startswith("```"):
            flush_para()
            flush_table()
            if in_code:
                for eline in code_buf:
                    if eline.strip():
                        add_equation(doc, equation_math(eline))
                code_buf = []
                in_code = False
            else:
                in_code = True
            continue
        if in_code:
            code_buf.append(line)
            continue

        # table rows
        if line.strip().startswith("|"):
            flush_para()
            table_buf.append(split_row(line))
            continue
        flush_table()

        if not line.strip():
            flush_para()
            continue

        if line.startswith("### "):
            flush_para()
            doc.add_heading(line[4:].strip(), level=2)
            continue
        if line.startswith("## "):
            flush_para()
            doc.add_heading(line[3:].strip(), level=1)
            continue

        if line.startswith("- "):
            flush_para()
            p = doc.add_paragraph(style="List Bullet")
            add_rich(p, line[2:].strip())
            continue

        para_buf.append(line.strip())

    flush_para()
    flush_table()


def main():
    doc = Document()

    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)

    tp = doc.add_paragraph()
    tp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = tp.add_run(TITLE)
    r.bold = True
    r.font.size = Pt(15)
    if SUBTITLE:
        sp = doc.add_paragraph()
        sp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        sr = sp.add_run(SUBTITLE)
        sr.italic = True
        sr.font.size = Pt(10)
        sr.font.color.rgb = RGBColor(0x66, 0x66, 0x66)

    for fname in FILES:
        process_file(doc, BASE / fname)

    out = BASE / "manuscript_draft.docx"
    doc.save(out)
    print(f"wrote {out.resolve()}  ({len(doc.paragraphs)} paragraphs, {len(doc.tables)} tables)")


if __name__ == "__main__":
    main()
